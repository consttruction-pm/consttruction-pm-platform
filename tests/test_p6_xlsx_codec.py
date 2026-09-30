from io import BytesIO
from openpyxl import Workbook
import pytest
from construction_pm.backend_p0.models import BackendScope
from construction_pm.p6_mapping_registry import P6MappingFormat
from construction_pm.p6_xlsx_codec import P6XlsxCodec,P6XlsxCodecError

def scope(): return BackendScope(tenant_id="t1",project_id="p1",project_revision=4)

def test_xlsx_round_trip():
    codec=P6XlsxCodec()
    from construction_pm.p6_interchange_mapping import P6InterchangeResult
    raw=codec.encode((P6InterchangeResult({"Id":"A-10","Name":"Foundation"},{"p6.xlsx.sheet":"Activities","vendor.custom":"keep"}),),scope())
    rows=codec.decode(raw,scope())
    assert rows[0].format is P6MappingFormat.XLSX
    assert rows[0].values=={"Id":"A-10","Name":"Foundation"}
    assert rows[0].extensions["vendor.custom"]=="keep"

def test_rejects_duplicate_headers():
    wb=Workbook(); ws=wb.active; ws.append(["Id","Id"]); ws.append(["A","B"])
    out=BytesIO(); wb.save(out)
    with pytest.raises(P6XlsxCodecError,match="DUPLICATE_HEADER:Sheet"): P6XlsxCodec().decode(out.getvalue(),scope())

def test_requires_sheet_metadata():
    from construction_pm.p6_interchange_mapping import P6InterchangeResult
    with pytest.raises(P6XlsxCodecError,match="MISSING_XLSX_SHEET"): P6XlsxCodec().encode((P6InterchangeResult({"Id":"A"},{}),),scope())

def test_rejects_oversized_sheet_name_instead_of_truncating():
    from construction_pm.p6_interchange_mapping import P6InterchangeResult
    sheet = "A" * 32
    with pytest.raises(P6XlsxCodecError, match="XLSX_SHEET_NAME_TOO_LONG"):
        P6XlsxCodec().encode((P6InterchangeResult({"Id": "A"}, {"p6.xlsx.sheet": sheet}),),scope())

def _extension_workbook(rows):
    wb=Workbook(); ws=wb.active; ws.title="Activities"; ws.append(["Id"]); ws.append(["A-10"])
    ext=wb.create_sheet("__ConstructionPMExtensions"); ext.sheet_state="hidden"
    ext.append(["sheet","key","value"])
    for row in rows: ext.append(row)
    out=BytesIO(); wb.save(out); return out.getvalue()

def test_rejects_malformed_extension_metadata():
    payload=_extension_workbook([("Activities",None,"lost")])
    with pytest.raises(P6XlsxCodecError,match="INVALID_EXTENSION_ROW:2"):
        P6XlsxCodec().decode(payload,scope())

def test_rejects_duplicate_extension_keys_instead_of_overwriting():
    payload=_extension_workbook([("Activities","vendor.custom","one"),("Activities","vendor.custom","two")])
    with pytest.raises(P6XlsxCodecError,match="DUPLICATE_EXTENSION_KEY:Activities:vendor.custom"):
        P6XlsxCodec().decode(payload,scope())

def test_rejects_extension_reference_to_missing_sheet():
    payload=_extension_workbook([("MissingActivities","vendor.custom","lost")])
    with pytest.raises(P6XlsxCodecError,match="EXTENSION_SHEET_NOT_FOUND:MissingActivities"):
        P6XlsxCodec().decode(payload,scope())
