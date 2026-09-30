from __future__ import annotations
from dataclasses import dataclass
from io import BytesIO
from typing import Any, Sequence
from openpyxl import Workbook, load_workbook
from .backend_p0.models import BackendScope
from .p6_interchange_mapping import P6InterchangeResult, P6InterchangeRow
from .p6_mapping_registry import P6MappingFormat

class P6XlsxCodecError(ValueError): pass
SHEET_EXTENSION="__ConstructionPMExtensions"
SHEET_EXTENSION_KEY="key"; SHEET_EXTENSION_VALUE="value"

@dataclass(frozen=True)
class P6XlsxCodec:
    format:P6MappingFormat=P6MappingFormat.XLSX

    def decode(self,payload:bytes,scope:BackendScope)->Sequence[P6InterchangeRow]:
        scope.validate()
        try: wb=load_workbook(BytesIO(payload),data_only=False)
        except Exception as exc: raise P6XlsxCodecError("INVALID_XLSX") from exc
        sheets=[s for s in wb.sheetnames if s!=SHEET_EXTENSION]
        if not sheets: raise P6XlsxCodecError("NO_XLSX_DATA_SHEET")
        ext=self._read_extensions(wb)
        rows=[]
        for sheet in sheets:
            ws=wb[sheet]
            headers=[c.value for c in ws[1]]
            if not headers or any(h is None or not str(h).strip() for h in headers): raise P6XlsxCodecError(f"INVALID_HEADERS:{sheet}")
            if len(set(headers))!=len(headers): raise P6XlsxCodecError(f"DUPLICATE_HEADER:{sheet}")
            for values in ws.iter_rows(min_row=2,values_only=True):
                if all(v is None for v in values): continue
                row_ext={"p6.xlsx.sheet":sheet}; row_ext.update(ext.get(sheet,{}))
                rows.append(P6InterchangeRow(scope=scope,format=self.format,values={str(h):v for h,v in zip(headers,values)},extensions=row_ext))
        return tuple(rows)

    def encode(self,rows:Sequence[P6InterchangeResult],scope:BackendScope)->bytes:
        scope.validate()
        wb=Workbook(); wb.remove(wb.active); extensions=[]
        grouped={}
        for row in rows:
            sheet=row.extensions.get("p6.xlsx.sheet")
            if not isinstance(sheet,str) or not sheet.strip(): raise P6XlsxCodecError("MISSING_XLSX_SHEET")
            if sheet==SHEET_EXTENSION: raise P6XlsxCodecError("RESERVED_XLSX_SHEET")
            grouped.setdefault(sheet,[]).append(row)
        for sheet,items in grouped.items():
            if len(sheet) > 31: raise P6XlsxCodecError("XLSX_SHEET_NAME_TOO_LONG")
            ws=wb.create_sheet(sheet); fields=[]
            for row in items:
                for f in row.values:
                    if f not in fields: fields.append(f)
                for k,v in row.extensions.items():
                    if k!="p6.xlsx.sheet": extensions.append((sheet,k,v))
            ws.append(fields)
            for row in items: ws.append([row.values.get(f) for f in fields])
        if extensions:
            ws=wb.create_sheet(SHEET_EXTENSION)
            ws.sheet_state="hidden"; ws.append(["sheet",SHEET_EXTENSION_KEY,SHEET_EXTENSION_VALUE])
            for sheet,key,value in sorted(extensions): ws.append([sheet,key,self._stringify(value)])
        out=BytesIO(); wb.save(out); return out.getvalue()

    @staticmethod
    def _read_extensions(wb):
        if SHEET_EXTENSION not in wb.sheetnames:return {}
        ws=wb[SHEET_EXTENSION]
        headers=[c.value for c in ws[1]]
        if headers != ["sheet",SHEET_EXTENSION_KEY,SHEET_EXTENSION_VALUE]:
            raise P6XlsxCodecError("INVALID_EXTENSION_HEADERS")
        result={}
        for row_number,(sheet,key,value) in enumerate(ws.iter_rows(min_row=2,max_col=3,values_only=True),2):
            if sheet is None and key is None and value is None:
                continue
            if not isinstance(sheet,str) or not sheet.strip() or not isinstance(key,str) or not key.strip():
                raise P6XlsxCodecError(f"INVALID_EXTENSION_ROW:{row_number}")
            sheet=str(sheet); key=str(key)
            if sheet==SHEET_EXTENSION:
                raise P6XlsxCodecError(f"INVALID_EXTENSION_ROW:{row_number}")
            bucket=result.setdefault(sheet,{})
            if key in bucket:
                raise P6XlsxCodecError(f"DUPLICATE_EXTENSION_KEY:{sheet}:{key}")
            bucket[key]=value
        return result

    @staticmethod
    def _stringify(value): return "" if value is None else str(value)

__all__=["P6XlsxCodec","P6XlsxCodecError"]
