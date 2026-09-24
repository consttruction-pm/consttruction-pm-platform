from __future__ import annotations
from io import BytesIO
from decimal import Decimal
from openpyxl import Workbook, load_workbook
from .models import Resource, ResourceAssignment
from .excel_schema import RESOURCE_COLUMNS, ASSIGNMENT_COLUMNS

SCHEMA_VERSION = "1.0"

def export_resources_xlsx(resources: list[Resource], assignments: list[ResourceAssignment]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resources"
    ws.append(["Schema Version", SCHEMA_VERSION])
    ws.append([c.name for c in RESOURCE_COLUMNS])
    for r in resources:
        ws.append([r.id, r.code, r.name, r.type.value, r.unit, r.calendar_id, r.active])
    wa = wb.create_sheet("Assignments")
    wa.append(["Schema Version", SCHEMA_VERSION])
    wa.append([c.name for c in ASSIGNMENT_COLUMNS])
    for a in assignments:
        wa.append([
            a.activity_id, a.resource_id, a.planned_units, a.actual_units,
            a.remaining_units, a.planned_cost, a.actual_cost, a.remaining_cost,
        ])
    output = BytesIO()
    wb.save(output)
    return output.getvalue()

def import_assignment_rows_xlsx(data: bytes) -> list[dict]:
    wb = load_workbook(BytesIO(data), data_only=False)
    if "Assignments" not in wb.sheetnames:
        raise ValueError("Assignments sheet is required")
    ws = wb["Assignments"]
    if ws["A1"].value != "Schema Version" or str(ws["B1"].value) != SCHEMA_VERSION:
        raise ValueError("Unsupported Excel schema version")
    headers = [cell.value for cell in ws[2]]
    expected = [c.name for c in ASSIGNMENT_COLUMNS]
    if headers != expected:
        raise ValueError("Invalid Assignments headers")
    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if all(value is None for value in row):
            continue
        rows.append({
            "activity_id": str(row[0]),
            "resource_id": str(row[1]),
            "planned_units": Decimal(str(row[2])),
            "actual_units": Decimal(str(row[3])),
            "remaining_units": Decimal(str(row[4])),
            "planned_cost": Decimal(str(row[5])),
            "actual_cost": Decimal(str(row[6])),
            "remaining_cost": Decimal(str(row[7])),
        })
    return rows
